import collections
import collections.abc
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

def create_poster():
    prs = Presentation()
    
    # Set slide size to A4 portrait
    prs.slide_width = Inches(8.27)
    prs.slide_height = Inches(11.69)
    
    # Use blank slide layout
    blank_slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_slide_layout)
    
    # Define colors
    navy_blue = RGBColor(0x1F, 0x38, 0x64)
    light_gray = RGBColor(0xF2, 0xF2, 0xF2)
    white = RGBColor(0xFF, 0xFF, 0xFF)
    dark_gray = RGBColor(0x40, 0x40, 0x40)

    # --- Header ---
    # Top banner (Navy Blue)
    shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(8.27), Inches(1.5))
    shape.fill.solid()
    shape.fill.fore_color.rgb = navy_blue
    shape.line.fill.background()
    
    # Header Text
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.2), Inches(7.27), Inches(1))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "GRADUATION PROJECT 2026\nFACULTY OF COMPUTER SCIENCE & AI"
    run.font.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = white
    
    # --- Content Sections ---
    sections = [
        {
            "title": "Problem",
            "content": "• Manual prescription handling leads to errors and delays.\n• Inefficient inventory management causes stockouts and expired medicines.\n• Lack of integrated software connecting digital orders to robotic hardware."
        },
        {
            "title": "Goal",
            "content": "Develop an integrated smart pharmacy ecosystem (APMS & Roshetety) to bridge patient-facing digital prescriptions with staff-facing AI inventory and robotic dispensing."
        },
        {
            "title": "Approach",
            "content": "• Build bilingual patient mobile app for OCR prescription scanning.\n• Create comprehensive staff dashboard for inventory and AI analytics.\n• Implement machine learning for demand forecasting and expiry prediction.\n• Integrate robotic dispensing arm via real-time WebSocket communication."
        },
        {
            "title": "Project Title",
            "content": "APMS & ROSHETETY: INTEGRATED SMART PHARMACY ECOSYSTEM"
        },
        {
            "title": "Technologies Used",
            "content": "Flutter (Dart), Node.js, Express, PostgreSQL, Python, XGBoost, LangChain, Keras, WebSockets, Serial Communication"
        },
        {
            "title": "Results and Conclusion",
            "content": "• Successfully integrated patient app, staff dashboard, and robotic hardware.\n• Achieved high accuracy in demand forecasting (XGBoost) and expiry prediction.\n• Reduced prescription processing time and minimized manual inventory tracking errors."
        }
    ]
    
    current_y = 1.8
    for sec in sections:
        # Title Box (Navy Blue)
        t_shape = slide.shapes.add_shape(1, Inches(0.5), Inches(current_y), Inches(3), Inches(0.4))
        t_shape.fill.solid()
        t_shape.fill.fore_color.rgb = navy_blue
        t_shape.line.fill.background()
        
        tf = t_shape.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.LEFT
        run = p.add_run()
        run.text = sec["title"]
        run.font.bold = True
        run.font.size = Pt(16)
        run.font.color.rgb = white
        
        # Content Box
        if sec["title"] == "Project Title":
             c_shape = slide.shapes.add_shape(1, Inches(0.5), Inches(current_y + 0.5), Inches(7.27), Inches(0.8))
             c_shape.fill.solid()
             c_shape.fill.fore_color.rgb = white
             c_shape.line.color.rgb = navy_blue
             c_shape.line.width = Pt(2)
             
             tf = c_shape.text_frame
             tf.word_wrap = True
             p = tf.paragraphs[0]
             p.alignment = PP_ALIGN.CENTER
             run = p.add_run()
             run.text = sec["content"]
             run.font.bold = True
             run.font.size = Pt(20)
             run.font.color.rgb = navy_blue
             current_y += 1.5
        else:
             c_shape = slide.shapes.add_shape(1, Inches(0.5), Inches(current_y + 0.5), Inches(4.5), Inches(1.0))
             c_shape.fill.solid()
             c_shape.fill.fore_color.rgb = white
             c_shape.line.color.rgb = light_gray
             
             tf = c_shape.text_frame
             tf.word_wrap = True
             p = tf.paragraphs[0]
             p.alignment = PP_ALIGN.LEFT
             run = p.add_run()
             run.text = sec["content"]
             run.font.size = Pt(12)
             run.font.color.rgb = dark_gray
             current_y += 1.6

    # --- Footer ---
    # Bottom banner (Navy Blue)
    f_shape = slide.shapes.add_shape(1, Inches(0), Inches(10.8), Inches(8.27), Inches(0.89))
    f_shape.fill.solid()
    f_shape.fill.fore_color.rgb = navy_blue
    f_shape.line.fill.background()
    
    # Students
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(10.9), Inches(3.5), Inches(0.7))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "STUDENT NAMES\nOmar Khaled, Sarah Mahmoud, Ahmed Ali"
    run.font.size = Pt(12)
    run.font.color.rgb = white

    # Supervisor
    txBox = slide.shapes.add_textbox(Inches(4.2), Inches(10.9), Inches(3.5), Inches(0.7))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    run = p.add_run()
    run.text = "SUPERVISED BY:\nProf. Dr. Ayman El-Ghazaly" # Placeholder based on image
    run.font.size = Pt(12)
    run.font.color.rgb = white

    prs.save("project_poster.pptx")

create_poster()
