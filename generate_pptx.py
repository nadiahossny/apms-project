import collections.abc
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt

prs = Presentation()

# Slide 1: Title Slide
slide_layout = prs.slide_layouts[0] # Title slide
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "roshetety"
subtitle.text = "Automated Pharmacy Management and Robotic Dispensing Ecosystem"

# Slide 2: Background & Problem Statement
slide_layout = prs.slide_layouts[1] # Title and Content
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Background & Problem Statement"
tf = slide.placeholders[1].text_frame
tf.text = "High incidence of critical errors in manual medication dispensing"
p = tf.add_paragraph()
p.text = "Inefficient reliance on legacy, paper-based operational workflows"
p = tf.add_paragraph()
p.text = "Prohibitive costs of existing enterprise robotic solutions"
p = tf.add_paragraph()
p.text = "Current robotics strictly limited to massive hospital infrastructures"

# Slide 3: Literature Review & Competitor Analysis
slide_layout = prs.slide_layouts[5] # Title Only
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Literature Review & Competitor Analysis"

rows = 6
cols = 4
left = Inches(0.5)
top = Inches(1.5)
width = Inches(9.0)
height = Inches(4.0)

table = slide.shapes.add_table(rows, cols, left, top, width, height).table

# Set column widths
table.columns[0].width = Inches(2.5)
table.columns[1].width = Inches(2.0)
table.columns[2].width = Inches(2.0)
table.columns[3].width = Inches(2.5)

# Headers
headers = ["Features", "Software-Only Apps", "Enterprise Robotics", "roshetety (Proposed)"]
for i, header in enumerate(headers):
    table.cell(0, i).text = header

data = [
    ["Automated Physical Dispensing", "No", "Yes", "Yes"],
    ["Eliminates Manual Data Entry", "No (Heavy manual entry)", "Yes", "Yes"],
    ["Social/Patient App for Prescriptions", "Yes", "No", "Yes"],
    ["Real-Time User Order Tracking", "Yes", "No", "Yes"],
    ["Affordable/Accessible", "Yes", "No (Expensive)", "Yes"]
]

for row_idx, row_data in enumerate(data):
    for col_idx, cell_data in enumerate(row_data):
        table.cell(row_idx + 1, col_idx).text = cell_data

# Slide 4: User Analysis & Stakeholder Needs
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "User Analysis & Stakeholder Needs"
tf = slide.placeholders[1].text_frame
tf.text = "[Insert Survey Results / Charts Here]"
p = tf.add_paragraph()
p.text = "Pharmacist Needs:"
p.level = 0
p = tf.add_paragraph()
p.text = "Significant reduction in manual data entry requirements"
p.level = 1
p = tf.add_paragraph()
p.text = "Mitigation of cognitive and physical fatigue during high-volume periods"
p.level = 1
p = tf.add_paragraph()
p.text = "Patient Needs:"
p.level = 0
p = tf.add_paragraph()
p.text = "Transparent, real-time order status tracking"
p.level = 1
p = tf.add_paragraph()
p.text = "Minimization of on-site prescription fulfillment wait times"
p.level = 1

# Slide 5: HCI & Design Requirements
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "HCI & Design Requirements"
tf = slide.placeholders[1].text_frame
tf.text = "Minimalist Interfaces: Streamlined data presentation to reduce visual clutter"
p = tf.add_paragraph()
p.text = "Accessible Interaction Targets: Large buttons to accommodate rapid, precise input"
p = tf.add_paragraph()
p.text = "Color-Coded Alert Systems: Instant visual prioritization of system states"
p = tf.add_paragraph()
p.text = "Primary Objective: Strict reduction of cognitive load and operational fatigue for pharmacists"

# Slide 6: The Proposed Solution
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "The Proposed Solution"
tf = slide.placeholders[1].text_frame
tf.text = "Automated Dispensing Robotic Arm: Hardware execution of medication retrieval"
p = tf.add_paragraph()
p.text = "Pharmacy Management Dashboard: Centralized system administration and oversight"
p = tf.add_paragraph()
p.text = "Patient-Facing Application: Interfaces for prescription submission and order management"

# Slide 7: System Architecture & Workflow
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "System Architecture & Workflow"
tf = slide.placeholders[1].text_frame
tf.text = "Sequential Operation Flow:"
p = tf.add_paragraph()
p.text = "Patient submits request via mobile application"
p.level = 1
p = tf.add_paragraph()
p.text = "System backend processes and validates prescription data"
p.level = 1
p = tf.add_paragraph()
p.text = "Management dashboard receives and queues the verified order"
p.level = 1
p = tf.add_paragraph()
p.text = "Robotic arm executes precise physical dispensing protocol"
p.level = 1
p = tf.add_paragraph()
p.text = "Application updates and tracks status in real-time"
p.level = 1
p = tf.add_paragraph()
p.text = "Technical Integration:"
p.level = 0
p = tf.add_paragraph()
p.text = "Frontend interfaces: React.js"
p.level = 1
p = tf.add_paragraph()
p.text = "Backend processing: Node.js"
p.level = 1
p = tf.add_paragraph()
p.text = "Relational Database: PostgreSQL"
p.level = 1
p = tf.add_paragraph()
p.text = "Direct Hardware-Software synchronization"
p.level = 1

# Slide 8: Full-Stack Implementation & Hardware
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Full-Stack Implementation & Hardware"
tf = slide.placeholders[1].text_frame
tf.text = "System Realization:"
p = tf.add_paragraph()
p.text = "Deployment of the React.js and Node.js micro-architecture"
p.level = 1
p = tf.add_paragraph()
p.text = "Integration of PostgreSQL for robust transactional data management"
p.level = 1
p = tf.add_paragraph()
p.text = "Calibration and deployment of the automated dispensing hardware"
p.level = 1
p = tf.add_paragraph()
p.text = "[Insert UI Screenshots]"
p.level = 0
p = tf.add_paragraph()
p.text = "[Insert Robotic Arm Photo]"
p.level = 0

# Slide 9: Evaluation & Results
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "Evaluation & Results"
tf = slide.placeholders[1].text_frame
tf.text = "Dispensing Accuracy: Error rate measurement versus manual baseline"
p = tf.add_paragraph()
p.text = "Data Entry Efficiency: Percentage reduction in manual input time"
p = tf.add_paragraph()
p.text = "System Latency: End-to-end response time from app request to hardware execution"

# Slide 10: References
slide_layout = prs.slide_layouts[1]
slide = prs.slides.add_slide(slide_layout)
slide.shapes.title.text = "References"
tf = slide.placeholders[1].text_frame
tf.text = "[Insert Competitor Application URLs]"
p = tf.add_paragraph()
p.text = "[Insert Referenced CS/HCI Literature]"

prs.save("roshetety_presentation.pptx")
