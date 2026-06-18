import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys
import re
import os

def create_docx():
    doc = docx.Document()
    
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)
    
    md_path = 'd:/march-26/apms-project/thesis_contribution.md'
    
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    for line in lines:
        line = line.strip()
        if not line:
            doc.add_paragraph()
            continue
            
        # Handle Images
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)$', line)
        if img_match:
            alt_text = img_match.group(1)
            img_path = img_match.group(2)
            
            # Clean path
            if img_path.startswith('file:///'):
                img_path = img_path[8:]
            
            # Use absolute path assuming images are in d:/march-26/apms-project/
            img_name = os.path.basename(img_path)
            local_img_path = os.path.join('d:/march-26/apms-project/', img_name)
            
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                if os.path.exists(local_img_path):
                    r = p.add_run()
                    r.add_picture(local_img_path, width=Inches(6.0))
                    # Add caption
                    caption = doc.add_paragraph()
                    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    c_run = caption.add_run(f"Figure: {alt_text}")
                    c_run.font.name = 'Times New Roman'
                    c_run.font.size = Pt(10)
                    c_run.font.italic = True
                else:
                    r = p.add_run(f"[Image not found: {local_img_path}]")
            except Exception as e:
                r = p.add_run(f"[Error loading image: {e}]")
            continue
            
        # Handle Headings
        if line.startswith('# CHAPTER'):
            p = doc.add_paragraph()
            run = p.add_run(line.replace('# ', ''))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.bold = True
        elif line.startswith('## '):
            p = doc.add_paragraph()
            run = p.add_run(line.replace('## ', ''))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
        elif line.startswith('### '):
            p = doc.add_paragraph()
            run = p.add_run(line.replace('### ', ''))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.underline = True
        elif line.startswith('<u>') and line.endswith('</u>'):
            p = doc.add_paragraph()
            run = p.add_run(line.replace('<u>', '').replace('</u>', ''))
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.underline = True
        elif line.startswith('* ') or line.startswith('- '):
            p = doc.add_paragraph(style='List Bullet')
            text = line[2:]
            _process_inline(p, text)
        elif re.match(r'^\d+\.\s*', line):
            p = doc.add_paragraph(style='List Number')
            content = re.sub(r'^\d+\.\s*', '', line)
            _process_inline(p, content)
        else:
            p = doc.add_paragraph()
            _process_inline(p, line)
            
    doc.save('thesis_contribution.docx')

def _process_inline(paragraph, text):
    """Basic inline bold handling for normal text"""
    if '**' in text:
        parts = text.split('**')
        for i, part in enumerate(parts):
            if part:
                r = paragraph.add_run(part)
                r.font.name = 'Times New Roman'
                r.font.size = Pt(12)
                if i % 2 == 1:
                    r.font.bold = True
    else:
        r = paragraph.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)

create_docx()
